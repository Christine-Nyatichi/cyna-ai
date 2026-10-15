import re
from datetime import datetime, timezone
from io import BytesIO

from docx import Document as DocxDocument
from fpdf import FPDF

DOCX_MIMETYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _clean_inline(text):
    """Removes Markdown marks like **bold**, `code` and [links](url)."""
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return re.sub(r"\*\*|__|`", "", text).strip()


def parse_notes(text):
    """Turns notes into a list of (kind, text, level) blocks.

    kind is one of: h1, h2, h3, bullet, number, para, blank
    """
    blocks = []
    for raw in text.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            continue

        if not stripped:
            if blocks and blocks[-1][0] != "blank":
                blocks.append(("blank", "", 0))
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            blocks.append((f"h{level}", _clean_inline(heading.group(2)), 0))
            continue

        bullet = re.match(r"^[-*\u2022]\s+(.*)$", stripped)
        if bullet:
            indent = len(line) - len(line.lstrip())
            blocks.append(("bullet", _clean_inline(bullet.group(1)), 1 if indent >= 2 else 0))
            continue

        number = re.match(r"^(\d+)[.)]\s+(.*)$", stripped)
        if number:
            blocks.append(("number", f"{number.group(1)}. {_clean_inline(number.group(2))}", 0))
            continue

        blocks.append(("para", _clean_inline(stripped), 0))

    return blocks


def safe_filename(title):
    name = re.sub(r"[^A-Za-z0-9_-]+", "-", title or "").strip("-")[:40]
    return name or "notes"


def _today():
    return datetime.now(timezone.utc).strftime("%B %d, %Y")


def _latin(text):
    """The standard PDF fonts only know Latin-1, so swap common symbols and replace the rest."""
    swaps = {
        "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
        "\u2013": "-", "\u2014": "-", "\u2022": "-", "\u2026": "...", "\u00a0": " ",
    }
    for old, new in swaps.items():
        text = text.replace(old, new)
    return text.encode("latin-1", "replace").decode("latin-1")


def build_docx(title, blocks):
    doc = DocxDocument()
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"Created with CYNA AI on {_today()}")

    for kind, text, level in blocks:
        if kind == "blank" or not text:
            continue
        if kind in ("h1", "h2", "h3"):
            doc.add_heading(text, level=int(kind[1]))
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet 2" if level else "List Bullet")
        else:
            doc.add_paragraph(text)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_pdf(title, blocks):
    pdf = FPDF()
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    def write(text, size, style="", height=6, indent=0, space_before=0):
        if space_before:
            pdf.ln(space_before)
        pdf.set_font("Helvetica", style, size)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(0, height, _latin(text), new_x="LMARGIN", new_y="NEXT")

    write(title, 22, "B", 10)
    write(f"Created with CYNA AI on {_today()}", 9, "", 5)
    pdf.ln(4)

    for kind, text, level in blocks:
        if kind == "blank":
            pdf.ln(3)
            continue
        if not text:
            continue
        if kind == "h1":
            write(text, 17, "B", 8, space_before=3)
        elif kind == "h2":
            write(text, 14, "B", 7, space_before=2)
        elif kind == "h3":
            write(text, 12, "B", 6, space_before=1)
        elif kind == "bullet":
            write("- " + text, 11, "", 6, indent=6 + 6 * level)
        elif kind == "number":
            write(text, 11, "", 6, indent=6)
        else:
            write(text, 11, "", 6)

    return bytes(pdf.output())